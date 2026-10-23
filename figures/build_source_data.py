"""Build data/figure_source_data.xlsx (plotted values of every figure, one sheet per figure)
from the model and the Monte Carlo results. Sheets that hold literature data or a schematic
(Fig1, FigA.1, FigA.2) are copied unchanged from the existing workbook.
Run from the repository root after model/monte_carlo.py:  python figures/build_source_data.py
"""
import csv, os
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment
from common import DATA, load_mc
from hf_tea_lca import P, GRADES, ROUTES, totals, flat_breakdown, with_, SENS, fsa_upstream, fsa_by_share

OUT = os.path.join(DATA, 'figure_source_data.xlsx')
old = openpyxl.load_workbook(OUT)
wb = openpyxl.Workbook(); wb.remove(wb.active)
GL = {'Technical': 'Technical', 'HighPurity': 'High-purity', 'Electronic': 'Electronic'}
COLS = [(r, g) for g in GRADES for r in ROUTES]
B = Font(bold=True)

def sheet(name, title, note=None):
    ws = wb.create_sheet(name); ws.append([title]); ws['A1'].font = B
    if note: ws.append([note])
    ws.append([]); return ws
def header(ws, row):
    ws.append(row)
    for c in ws[ws.max_row]: c.font = B
def copy_sheet(name, new=None):
    src = old[name] if name in old.sheetnames else old[new]; ws = wb.create_sheet(new or name)
    for row in src.iter_rows(values_only=True): ws.append(list(row))
    if new and isinstance(ws['A1'].value, str): ws['A1'] = ws['A1'].value.replace('Fig. S2', 'Fig. A.2')
    ws['A1'].font = B

# README
ws = sheet('README', 'Fluorosilicic acid as an alternative to fluorspar for electronic-grade HF production: source data of the figures (sheet name = figure number)',
           'Model: model/hf_tea_lca.py · Monte Carlo: model/monte_carlo.py -> results/mc_results.npz (N = 5,000, seed 20260929) · Built by figures/build_source_data.py')
header(ws, ['Sheet', 'Figure', 'Script', 'Source of values'])
for r in [('Fig1', 'Overview of the modeled system and research questions', '- (schematic drawn by the authors)', 'No model values'),
          ('Fig2a', 'Unit production cost by route and grade', 'figures/fig2_cost_ghg.py', 'hf_tea_lca.flat_breakdown'),
          ('Fig2b', 'Cradle-to-gate GHG emissions by route and grade, with on-site emissions', 'figures/fig2_cost_ghg.py', 'hf_tea_lca.flat_breakdown (CtG) and totals (GtG)'),
          ('Fig3', 'Resource conservation and its emission trade-off', 'figures/fig3_resource.py', 'hf_tea_lca scenarios; USGS MCS 2026; MarketsandMarkets 2025'),
          ('Fig4ab', 'One-at-a-time sensitivity with input values, and Monte Carlo cost gap', 'figures/fig4_sensitivity_qualification.py', 'hf_tea_lca.SENS; results/mc_results.npz'),
          ('Fig4c', 'One-time qualification cost absorbed', 'figures/fig4_sensitivity_qualification.py', 'Electronic-grade cost gap; Monte Carlo 5th-95th percentile'),
          ('Fig4b_draws', 'All Monte Carlo draws: cost and emissions by route and grade', 'model/monte_carlo.py', 'results/mc_results.npz'),
          ('Fig4b_inputs', 'Input values sampled in each Monte Carlo draw', 'model/monte_carlo.py', 'results/mc_results.npz'),
          ('FigA.1', 'Appendix A: HF price by grade and fluorspar production', 'figures/figA1_price_supply.py', 'Price sources in SI Table S3 and the source dossier; USGS MCS 2026'),
          ('FigS1', 'System boundary and multifunctional points', 'figures/figS1_system_boundary.py', 'SI Table S6'),
          ('FigA.2', 'Appendix A: producer concentration versus recovery of phosphate fluorine', '- (values below)', 'USGS MCS 2026')]:
    ws.append(list(r))
ws.append([]); ws.append(['The published figures were assembled and lettered in PowerPoint from the script output; plotted values are identical.'])
copy_sheet('Fig1')

# Fig2a / Fig2b
src = open(os.path.join(os.path.dirname(__file__), 'fig2_cost_ghg.py'), encoding='utf-8').read()
cm2 = eval(src.split('cmap2 = ')[1].split('\ncol2')[0]); col2 = eval(src.split('col2 = ')[1].split('\nfig,')[0])
cm3 = eval(src.split('cmap3 = ')[1].split('\ncol3')[0]); col3 = eval(src.split('col3 = ')[1].split('\nfig,')[0])
def agg(d, cmap):
    out = {k: 0.0 for k in dict.fromkeys(cmap.values())}
    for k, v in d.items():
        base = k.split(':')[-1]; purif = k.startswith('purif:')
        out[cmap.get(('purif' if purif else '') + base) or cmap.get(base)] += v
    return out
for name, title, idx, cmap, cols, scale, unit in (('Fig2a', 'Fig. 2a — Unit production cost by route and grade', 0, cm2, col2, 1000, '$/kg HF'),
                                                  ('Fig2b', 'Fig. 2b — Cradle-to-gate GHG emissions by route and grade', 1, cm3, col3, 1, 'kgCO2eq/kg HF')):
    ws = sheet(name, title, f'Unit: {unit}. Base case: economic allocation of the FSA upstream burden, integrated site, silica returned.')
    header(ws, ['Category'] + [f'{r}, {GL[g]}' for r, g in COLS])
    data = {c: agg(flat_breakdown(*c)[idx], cmap) for c in COLS}
    for cat in cols: ws.append([cat] + [round(data[c][cat]/scale, 4) for c in COLS])
    tot = [totals(r, g)[idx] for r, g in COLS]
    ws.append(['Total (net)'] + [round(v, 4) for v in tot])
    if idx == 1: ws.append(['On-site emissions (diamond marker)'] + [round(totals(r, g, bd='GtG')[1], 4) for r, g in COLS])
    ws.append([]); header(ws, ['Bracket label (per grade)'] + [GL[g] for g in GRADES])
    ws.append([f'Gap, fluorspar minus FSA ({unit})'] + [round(tot[2*i] - tot[2*i+1], 4) for i in range(3)])
    ws.append(['Ratio, fluorspar to FSA (-)'] + [round(tot[2*i]/tot[2*i+1], 4) for i in range(3)])

# Fig3
E = 'Electronic'; SPAR = P['spar_t']
def dghg(p): return (totals('FSA', E, p)[1] - totals('Fluorspar', E, p)[1])/SPAR
base = dghg(P)
ws = sheet('Fig3', 'Fig. 3 — Resource conservation and emission trade-off',
           'ΔGHG per kg fluorspar conserved = (CtG FSA − CtG fluorspar) / 2.14 kg; identical across grades; linear between endpoints')
header(ws, ['Panel', 'Case', 'ΔGHG at 0%', 'ΔGHG at 100%', 'Break-even share (%)', 'Unit', 'Source'])
for panel, lab, kw, srcs in (('a', 'Marginal FSA, limestone neutralization avoided (−1.5 kgCO2eq/kg HF)', dict(fsa_alloc=0.0, cr_neutral=-1.5), 'SI Table S7'),
                             ('a', 'Marginal FSA, lime neutralization avoided (−2.3)', dict(fsa_alloc=0.0, cr_neutral=-2.3), 'SI Table S7'),
                             ('b', 'Silica sold, sodium silicate substituted (−0.75)', dict(cr_silica=-0.75, silica_sold=1), 'SI Table S7 (KEITI)'),
                             ('b', 'Silica sold, silica gel substituted (−1.3)', dict(cr_silica=-1.3, silica_sold=1), 'SI Table S7 (GREET)')):
    end = dghg(with_(**kw))
    ws.append([panel, lab, round(base, 4), round(end, 4), round(base/(base - end)*100, 1), 'kgCO2eq/kg fluorspar', srcs])
ws.append([]); header(ws, ['Panel', 'Band', 'Low (Mt AHF/yr)', 'High (Mt AHF/yr)', 'Fluorspar avoided low (Mt/yr)', 'Fluorspar avoided high (Mt/yr)', 'Source'])
FSA_HFEQ = P['fsa_t']*6*20.006/144.09
AIN = P['ahf_in']   # demand bands on the gross basis: AHF input per t of electronic-grade HF
for lab, lo, hi, s in (('AHF for electronic-grade HF demand, 2025', 0.3*AIN, 0.4*AIN, 'Demand 0.3–0.4 Mt HF × 1.39 t AHF/t electronic-grade HF; SI Section S14'),
                       ('AHF for projected electronic-grade HF demand, 2035', 0.3*1.06**10*AIN, 0.4*1.06**10*AIN, 'MarketsandMarkets 2025–2030 value growth 6.0%/yr extended to 2035, × 1.39 t AHF/t (SI Section S14)'),
                       ('Current FSA recovery, world (AHF supplied)', 0.3/FSA_HFEQ, 1.0/FSA_HFEQ, '0.3–1.0 Mt HF-eq, order-of-magnitude estimate (SI Section S14)'),
                       ('Potential FSA recovery (AHF supplied)', 0.5/FSA_HFEQ, 2.0/FSA_HFEQ, '0.5–2.0 Mt HF-eq, order-of-magnitude estimate (SI Section S14)')):
    ws.append(['c', lab, round(lo, 3), round(hi, 3), round(lo*SPAR, 3), round(hi*SPAR, 3), s + '; 2.14 t fluorspar/t AHF'])
fl, ph = {}, {}
with open(os.path.join(DATA, 'usgs_mcs2026_production_2025.csv'), encoding='utf-8') as f:
    for r in csv.DictReader(f): fl[r['country']] = float(r['fluorspar_kt']); ph[r['country']] = float(r['phosphate_rock_kt'])
ws.append([]); header(ws, ['Panel', 'Commodity', 'China %', 'Second', 'Third', 'HHI', 'Source'])
for lab, d, tot in (('Fluorspar', fl, 10000), ('Phosphate rock', ph, 250000)):
    top = sorted(d.items(), key=lambda kv: -kv[1])[:3]
    ws.append(['d', lab, round(top[0][1]/tot*100, 1), f'{top[1][0]} {top[1][1]/tot*100:.1f}', f'{top[2][0]} {top[2][1]/tot*100:.1f}',
               round(sum((v/tot*100)**2 for v in d.values())), 'USGS MCS 2026 (2025 estimates)'])

# Fig4ab
TAB = {'Electricity price': '0.108, 0.118, 0.128 $/kWh', 'Natural gas price': '14.3, 16.7, 20.2 $/GJ', 'Steam price': '12.8, 18.8, 23.5 $/GJ',
       'Sulfuric acid price': '32, 60, 125 $/t', 'FSA price': '105, 190, 246 $/t', 'Fluorspar price': '405, 470, 504 $/t',
       'AHF plant staffing': '1.44, 1.92, 2.40 persons/kt', 'Kiln fuel': '4.0, 5.0, 6.8 GJ/t', 'Fluorspar-route electricity': '150, 250, 556 kWh/t',
       'FSA-route steam': '12, 14, 16 GJ/t', 'FSA-route electricity': '220, 270, 315 kWh/t', 'FSA-route capital cost': '2,254, 2,337, 2,466 $/(t/yr)',
       'Fluorspar-route capital cost': '1,388, 2,234, 2,355 $/(t/yr)', 'Location factor': '0.956, 1.0, 1.84',
       'Purification capital cost': '3,416, 4,764, 6,612 $/(t/yr)', 'Purification AHF input': '1.14, 1.39, 1.39 t/t',
       'Industrial-grade HF price': '787, 1,060, 1,600 $/t', 'Labor cost': '20.8, 58.7, 64.0 k$/person', 'Oleum price ratio': '1.0, 1.1, 2.5',
       'Gypsum landfill cost': '3.5, 20, 70 $/t', 'Phosphate value': '1,300, 1,000, 900 $/t', 'Phosphorus removal increment': '0, 0, 0.094 $/kg',
       'Sulfuric acid emission factor': '0.045, 0.17, 0.17 kg/kg', 'Electricity emission factor': '0.495, 0.53, 0.53 kg/kWh', 'Capacity utilization': '100%, 100%, 85%'}
def gap(p): return totals('Fluorspar', E, p)[0] - totals('FSA', E, p)[0]
g0 = gap(P)
ws = sheet('Fig4ab', 'Fig. 4a–b — (a) one-at-a-time sensitivity of the electronic-grade cost gap with the low, base, and high input values (column C), (b) Monte Carlo cost gap by grade',
           f'Base electronic-grade cost gap = {g0:.4f} $/kg HF; panel (a) plots the change in cost gap, (gap − base)/base × 100 (columns H–I), for the 12 inputs with the largest swing')
header(ws, ['Panel', 'Input (figure label)', 'Low, base, high', 'Gap at low value ($/kg HF)', 'Gap at high value ($/kg HF)', 'Swing', 'Plotted in Fig. 4a',
            'Change in cost gap at low value (%)', 'Change in cost gap at high value (%)', 'Bar label, low', 'Bar label, high'])
rows = []
for n, lo, hi in SENS:
    a_ = gap(with_(**lo)) if lo else g0; b_ = gap(with_(**hi)) if hi else g0
    rows.append((n, a_, b_, max(abs(a_ - g0), abs(b_ - g0))))
rows.sort(key=lambda t: -t[3])
for i, (n, a_, b_, s_) in enumerate(rows):
    pl, ph = (a_ - g0)/g0*100, (b_ - g0)/g0*100
    lab = lambda v: 'no label' if abs(v) < 0.05 else (f'{v:+.0f}%' if abs(v) >= 1 else f'{v:+.1f}%').replace('-', '−')
    ws.append(['a', n, TAB[n], round(a_, 4), round(b_, 4), round(s_, 4), 'yes' if i < 12 and s_ > 1e-3 else 'no',
               round(pl, 2), round(ph, 2), lab(pl), lab(ph)])
mc = load_mc()
ws.append([]); header(ws, ['Panel', 'Grade', 'P5', 'Median', 'P95', 'Mean', 'SD', 'Share of draws with gap > 0'])
for g in GRADES:
    d = mc[f'U_S_{g}'] - mc[f'U_F_{g}']
    ws.append(['b', GL[g]] + [round(float(v), 4) for v in np.percentile(d, [5, 50, 95])] + [round(float(np.mean(d)), 4), round(float(np.std(d)), 4)]
              + [round(float(np.mean(d > 0)), 4)])

# Fig4c
gm = mc[f'U_S_{E}'] - mc[f'U_F_{E}']; p5, p95 = np.percentile(gm, [5, 95])
ws = sheet('Fig4c', 'Fig. 4c — One-time qualification cost absorbed ($ million) = cost gap × annual volume × contract term (undiscounted)',
           f'Cost gap {g0:.4f} $/kg HF (= $ million per kt); Monte Carlo 5–95% {p5:.4f}–{p95:.4f}; world electronic-grade demand 350 kt/yr for the share column')
header(ws, ['Annual volume (kt/yr)', 'Share of demand at 350 kt/yr (%)', '1.5 yr', '3 yr', '5 yr', '10 yr', '5 yr, P5 gap', '5 yr, P95 gap'])
for v in np.arange(0, 8.01, 0.5):
    ws.append([float(v), round(v/350*100, 3)] + [round(g0*v*t, 3) for t in (1.5, 3, 5, 10)] + [round(p5*v*5, 3), round(p95*v*5, 3)])
ws.append([]); header(ws, ['Marker', 'Volume (kt/yr)', 'Share of demand (%)', 'Calculation'])
for lab, v, calc in (('5-yr term absorbs $10 million', 10/(g0*5), '10 / (gap × 5)'), ('18-month term absorbs $10 million', 10/(g0*1.5), '10 / (gap × 1.5)'),
                     ('5-yr term at P5 gap', 10/(p5*5), '10 / (P5 gap × 5)')):
    ws.append([lab, round(v, 3), f'{v/400*100:.2f}–{v/300*100:.2f}', calc])

# Fig4b_draws, Fig4b_inputs: every Monte Carlo draw (results/mc_results.npz)
ws = sheet('Fig4b_draws', 'Fig. 4b — All 5,000 Monte Carlo draws: unit production cost ($/kg HF) and cradle-to-gate emissions (kgCO2eq/kg HF) by route and grade',
           'Draw i uses the input values in row i of sheet Fig4b_inputs. Gap = fluorspar minus FSA. Fig. 4b plots the gap columns.')
hd = ['Draw']
for g in GRADES:
    hd += [f'UPC fluorspar, {GL[g]}', f'UPC FSA, {GL[g]}', f'Cost gap, {GL[g]}']
for g in GRADES:
    hd += [f'CtG fluorspar, {GL[g]}', f'CtG FSA, {GL[g]}']
header(ws, hd)
n_mc = len(mc[f'U_S_{E}'])
for i in range(n_mc):
    r = [i + 1]
    for g in GRADES:
        us, uf = float(mc[f'U_S_{g}'][i]), float(mc[f'U_F_{g}'][i]); r += [round(us, 5), round(uf, 5), round(us - uf, 5)]
    for g in GRADES:
        r += [round(float(mc[f'N_S_{g}'][i]), 5), round(float(mc[f'N_F_{g}'][i]), 5)]
    ws.append(r)
with open(os.path.join(os.path.dirname(DATA), 'inputs', 'parameters.csv'), encoding='utf-8') as f:
    PAR = {}
    for row in csv.DictReader(f): PAR.setdefault(row['parameter'], row)
IN_KEYS = [k[3:] for k in mc.files if k.startswith('in_')]
def in_label(k):
    if k == 'p_rm_scale': return 'Phosphorus removal increment, fraction of the high case', '-', 0.0, 1.0, 'uniform'
    r = PAR[{'ci_h2so4_ctg': 'ci_h2so4', 'ci_elec_ctg': 'ci_elec'}.get(k, k)]
    lo, hi = r['mc_low'].replace('CtG ', ''), r['mc_high'].replace('CtG ', '')
    desc = r['description'] + (', cradle-to-gate' if k.endswith('_ctg') else '')
    return desc, r['unit'].replace('0 / ', ''), float(lo), float(hi), 'triangular'
ws = sheet('Fig4b_inputs', 'Fig. 4b — Input values sampled in each Monte Carlo draw (N = 5,000, seed 20260929; SI Table S8)',
           'Rows 4–9 describe each column; draws start at row 10. Capital cost of the FSA-route plant is multiplied by the location factor in the model.')
lab = [in_label(k) for k in IN_KEYS]
for name, j in (('Parameter', None), ('Description', 0), ('Unit', 1), ('Low', 2), ('High', 3), ('Distribution', 4)):
    header(ws, [name] + (IN_KEYS if j is None else [l[j] for l in lab]))
for i in range(n_mc):
    ws.append([i + 1] + [round(float(mc['in_' + k][i]), 6) for k in IN_KEYS])
copy_sheet('FigA.1')
ws = sheet('FigS1', 'Fig. S1 — System boundary and multifunctional points (diagram, no plotted values)', 'See SI Table S6.')
header(ws, ['Point', 'Flow', 'Route', 'Base method', 'Base value'])
for r in (('M1', 'Upstream burden of FSA from the phosphoric acid plant', 'FSA', 'Economic allocation', f'{fsa_upstream(P):.2f} kgCO2eq/kg HF'),
          ('M2', 'Silica', 'FSA', 'No allocation, returned to the phosphoric acid plant', '0'),
          ('M3', 'Diluted sulfuric acid', 'FSA', 'No allocation, closed loop on an integrated site', '0'),
          ('M4', 'Fluorogypsum', 'Fluorspar', 'Economic allocation (value 0 $/t)', f"Landfill of unsold share +{P['gypsum_t']*(1-P['gyp_sold'])*P['ci_landfill']*(1-fsa_by_share(P)):.3f} kgCO2eq/kg HF"),
          ('M5', 'FSA co-product', 'Fluorspar', 'Economic allocation', f'{fsa_by_share(P)*100:.2f}% of burden'),
          ('M6', 'Industrial-grade HF from purification', 'Both', 'Physical allocation (HF mass)', 'Product carries 1.0 t AHF')):
    ws.append(list(r))
copy_sheet('FigS2', 'FigA.2')
for ws in wb.worksheets:
    ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 40
    for c in 'CDEFGH': ws.column_dimensions[c].width = 16
wb.save(OUT); print('saved', OUT)
